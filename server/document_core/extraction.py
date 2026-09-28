"""Text extraction shared by the standalone MCP and legacy Gateway."""

import codecs
from pathlib import Path
from .limits import MAX_PAGES, MAX_TEXT_CHARS, DocumentLimitError, validate_input, validate_workbook


def extract_text(path: Path, name: str) -> str:
    validate_input(path, name)
    suffix = Path(name).suffix.lower()
    blocks = []
    count = 0

    def add(text):
        nonlocal count
        text = str(text)
        remaining = MAX_TEXT_CHARS - count
        if remaining <= 0:
            return False
        blocks.append(text[:remaining])
        count += len(blocks[-1]) + 1
        return count < MAX_TEXT_CHARS

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if len(reader.pages) > MAX_PAGES:
            raise DocumentLimitError("PDF 页数超过处理上限。")
        for page in reader.pages:
            if not add(page.extract_text() or ""):
                break
    elif suffix == ".docx":
        from docx import Document

        document = Document(path)
        for paragraph in document.paragraphs:
            if not add(paragraph.text):
                break
        for table in document.tables:
            for row in table.rows:
                if not add("\t".join(cell.text for cell in row.cells)):
                    return "\n".join(blocks)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            validate_workbook(workbook)
            for sheet in workbook.worksheets:
                add(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if not add("\t".join("" if v is None else str(v) for v in row)):
                        return "\n".join(blocks)
        finally:
            workbook.close()
    elif suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(path)
        if len(deck.slides) > MAX_PAGES:
            raise DocumentLimitError("幻灯片数量超过处理上限。")
        for i, slide in enumerate(deck.slides, start=1):
            add(f"## Slide {i}")
            for shape in slide.shapes:
                if getattr(shape, "text", "") and not add(shape.text):
                    return "\n".join(blocks)
    else:
        # Bound the read itself; do not allocate the complete large input and
        # only then truncate it for a response.
        with path.open("rb") as source:
            raw = source.read(MAX_TEXT_CHARS * 4)
        for encoding in ("utf-8-sig", "gb18030", "latin-1"):
            try:
                # A byte budget may end inside a multi-byte character. Decode
                # complete characters rather than misclassifying valid UTF-8
                # as a different encoding solely because its tail was cut.
                text = codecs.getincrementaldecoder(encoding)().decode(raw, final=False)
                break
            except UnicodeDecodeError:
                continue
        add(text)
    return "\n".join(blocks)
