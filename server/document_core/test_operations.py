from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from openpyxl import Workbook, load_workbook
import pytest

from document_core.operations import edit_document_path


def _edit_docx(tmp_path, document, find, replace, all_occ=True):
    source, target = tmp_path / "source.docx", tmp_path / "edited.docx"
    document.save(source)
    edit_document_path(
        source,
        target,
        source.name,
        {
            "schema_version": 1,
            "format": "docx",
            "ops": [
                {
                    "op": "replace_text",
                    "find": find,
                    "replace": replace,
                    "all": all_occ,
                }
            ],
        },
    )
    return Document(target)


def _hyperlink(paragraph, text):
    relation = paragraph.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    run.text = text
    link.append(run)
    paragraph._p.append(link)


def test_docx_replacement_preserves_unmatched_runs_and_drawing(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Prefix ").bold = True
    run = p.add_run("old")
    run.italic = True
    drawing = OxmlElement("w:drawing")
    run._r.append(drawing)
    p.add_run(" suffix").underline = True
    result = _edit_docx(tmp_path, doc, "old", "new").paragraphs[0]
    assert [(r.text, r.bold, r.italic, r.underline) for r in result.runs] == [
        ("Prefix ", True, None, None),
        ("new", None, True, None),
        (" suffix", None, None, True),
    ]
    assert len(result.runs[1]._r.xpath("./w:drawing")) == 1


@pytest.mark.parametrize(
    "find,replace,expected,link_text",
    [
        ("old", "new", "new LINK suffix", "LINK"),
        ("LINK", "new link", "old new link suffix", "new link"),
        ("ld LI", "X", "oXNK suffix", "NK"),
    ],
)
def test_docx_replacement_preserves_hyperlinks(tmp_path, find, replace, expected, link_text):
    doc = Document()
    p = doc.add_paragraph("old ")
    _hyperlink(p, "LINK")
    p.add_run(" suffix")
    result = _edit_docx(tmp_path, doc, find, replace).paragraphs[0]
    assert result.text == expected
    assert result.hyperlinks[0].text == link_text
    assert result.hyperlinks[0].url == "https://example.com"


def test_docx_multiple_cross_run_matches_preserve_suffix_and_insert_controls(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("pre ab").bold = True
    p.add_run("c abc a").italic = True
    p.add_run("bc end").underline = True
    result = _edit_docx(tmp_path, doc, "abc", "X\tY\nZ").paragraphs[0]
    assert result.text == "pre X\tY\nZ X\tY\nZ X\tY\nZ end"
    assert result.runs[0].bold and result.runs[1].italic and result.runs[2].underline
    assert result.runs[2].text == " end"


def test_docx_replace_once_and_delete_text(tmp_path):
    doc = Document()
    doc.add_paragraph("old old")
    doc.add_paragraph("old")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "old"
    result = _edit_docx(tmp_path, doc, "old", "", all_occ=False)
    assert [p.text for p in result.paragraphs] == [" old", "old"]
    assert result.tables[0].cell(0, 0).text == "old"


def test_set_range_clears_cells_and_keeps_surrounding_content(tmp_path):
    source, target = tmp_path / "input.xlsx", tmp_path / "out.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["old", "untouched", "right"])
    ws.append(["old2", "old3"])
    wb.save(source)
    edit_document_path(
        source,
        target,
        source.name,
        {
            "schema_version": 1,
            "format": "xlsx",
            "ops": [
                {
                    "op": "set_range",
                    "start": "A1",
                    "values": [[None], [None, "new"]],
                }
            ],
        },
    )
    result = load_workbook(target)
    assert list(result.active.values) == [(None, "untouched", "right"), (None, "new", None)]
    result.close()


@pytest.mark.parametrize(
    "requested,expected",
    [
        ("result.pdf", "result.txt"),
        ("result", "result.txt"),
        ("result.TXT", "result.TXT"),
        (None, "source_edited.txt"),
    ],
)
def test_edit_output_name_matches_source_format(tmp_path, requested, expected):
    source, target = tmp_path / "source.txt", tmp_path / "target.txt"
    source.write_text("old", encoding="utf-8")
    output = edit_document_path(
        source,
        target,
        source.name,
        {
            "schema_version": 1,
            "format": "txt",
            "output_filename": requested,
            "ops": [{"op": "set_text", "text": "new"}],
        },
    )
    assert output.filename == expected
    assert output.media_type == "text/plain; charset=utf-8"
    assert target.read_text(encoding="utf-8") == "new"
